"""Small synthetic documents generated in memory; no operational files required."""
from io import BytesIO

TEXT='Blue team delivers 200 supply kits to the North District depot.'


def docx_bytes(text=TEXT):
    from docx import Document
    d=Document();d.add_paragraph(text)
    table=d.add_table(rows=2,cols=2)
    table.cell(0,0).text='Resource';table.cell(0,1).text='Quantity'
    table.cell(1,0).text='Water';table.cell(1,1).text='50'
    out=BytesIO();d.save(out);return out.getvalue()


def pptx_bytes():
    from pptx import Presentation
    from pptx.util import Inches
    d=Presentation();slide=d.slides.add_slide(d.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2)).text=TEXT
    table=slide.shapes.add_table(2,2,Inches(1),Inches(3),Inches(5),Inches(2)).table
    table.cell(0,0).text='Route';table.cell(0,1).text='Status'
    table.cell(1,0).text='North';table.cell(1,1).text='Open'
    slide.notes_slide.notes_text_frame.text='Confirm depot access before delivery.'
    out=BytesIO();d.save(out);return out.getvalue()


def pdf_bytes():
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    writer=PdfWriter();page=writer.add_blank_page(width=612,height=792)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream=DecodedStreamObject();stream.set_data(('BT /F1 12 Tf 50 700 Td ('+TEXT+') Tj ET').encode())
    page[NameObject('/Contents')]=writer._add_object(stream)
    out=BytesIO();writer.write(out);return out.getvalue()


def scanned_pdf_bytes():
    from PIL import Image,ImageDraw,ImageFont
    image=Image.new('RGB',(1500,600),'white')
    draw=ImageDraw.Draw(image)
    try:font=ImageFont.truetype('DejaVuSans.ttf',35)
    except OSError:font=ImageFont.load_default(size=35)
    draw.text((70,150),'Blue team delivers supply kits to the depot.',font=font,fill='black')
    out=BytesIO();image.save(out,format='PDF',resolution=150);return out.getvalue()
