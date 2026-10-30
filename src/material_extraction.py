"""Local document extraction. No document content is sent to external services."""
from io import BytesIO
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_UNITS = 5000
MIME = {'.pdf':'application/pdf', '.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document', '.pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation'}


def validate_container(data, filename):
    ext = Path(filename).suffix.lower()
    if ext not in MIME:
        raise ValueError('Use PDF, DOCX, or PPTX. Save legacy DOC/PPT files in the modern format first.')
    if not 0 < len(data) <= MAX_BYTES:
        raise ValueError('Document must be between 1 byte and 20 MB.')
    if ext == '.pdf':
        if not data.startswith(b'%PDF-'):
            raise ValueError('File does not contain a PDF header.')
    else:
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > 5000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024:
                    raise ValueError('Office document exceeds the expanded-size limit (80 MB / 5000 entries).')
                expected = 'word/document.xml' if ext == '.docx' else 'ppt/presentation.xml'
                if expected not in archive.namelist():
                    raise ValueError('Office document contents do not match its extension.')
                if any(i.flag_bits & 1 for i in infos):
                    raise ValueError('Encrypted Office files are not supported.')
        except zipfile.BadZipFile as exc:
            raise ValueError('Invalid Office document archive.') from exc
    return ext


def chunks_from_units(units):
    chunks, total = [], 0
    for locator, text, method in units:
        text = re.sub(r'[\t \u00a0]+', ' ', text.replace('\x00','')).strip()
        total += len(text)
        if total > MAX_TEXT:
            raise ValueError('Extracted text exceeds the 2 million character limit.')
        for start in range(0, len(text), 1800):
            part = text[start:start+1800]
            if part.strip():
                chunks.append({'text':part, 'locator':{**locator,'char_start':start,'char_end':start+len(part)}, 'method':method})
                if len(chunks) > MAX_UNITS:
                    raise ValueError('Document exceeds the 5000-passage limit.')
    if not chunks:
        raise ValueError('No readable text found. Image-only Word/PowerPoint content must be converted to PDF for OCR.')
    return chunks


def pdf_units(data):
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise ValueError('Encrypted PDFs are not supported. Upload an unencrypted copy.')
    if len(reader.pages) > 200:
        raise ValueError('PDF exceeds 200 pages. Split it into smaller documents.')
    with tempfile.TemporaryDirectory() as temp:
        source = Path(temp)/'source.pdf'
        source.write_bytes(data)
        for page_number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            method = 'pdf-text'
            if not text.strip():
                if not shutil.which('pdftoppm') or not shutil.which('tesseract'):
                    raise ValueError('Scanned PDF requires Poppler and Tesseract (included in Docker).')
                output = Path(temp)/'page'
                subprocess.run(['pdftoppm','-f',str(page_number),'-l',str(page_number),'-scale-to','2400','-png','-singlefile',str(source),str(output)],check=True,capture_output=True,timeout=30)
                text = subprocess.run(['tesseract',str(output)+'.png','stdout','-l','eng'],check=True,capture_output=True,text=True,timeout=30).stdout
                method = 'pdf-ocr-eng'
            yield {'kind':'page','page':page_number}, text, method


def docx_units(data):
    from docx import Document
    from docx.text.paragraph import Paragraph
    document = Document(BytesIO(data))
    paragraph, table = 0, 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph += 1
            yield {'kind':'paragraph','paragraph':paragraph}, block.text, 'docx-text'
        else:
            table += 1
            for number, row in enumerate(block.rows, 1):
                yield {'kind':'table_row','table':table,'row':number}, ' | '.join(cell.text for cell in row.cells), 'docx-table'
    for number, section in enumerate(document.sections, 1):
        for kind, part in [('header',section.header),('footer',section.footer)]:
            for index, paragraph in enumerate(part.paragraphs, 1):
                yield {'kind':kind,'section':number,'paragraph':index}, paragraph.text, 'docx-text'


def pptx_units(data):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    deck = Presentation(BytesIO(data))
    if len(deck.slides) > 300:
        raise ValueError('Presentation exceeds 300 slides.')
    def shapes_text(shapes, slide, path=''):
        for number, shape in enumerate(shapes, 1):
            shape_path = f'{path}{number}'
            locator = {'kind':'slide','slide':slide,'shape':shape_path}
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes_text(shape.shapes,slide,shape_path+'.')
            elif shape.has_table:
                for row_number, row in enumerate(shape.table.rows, 1):
                    yield {**locator,'row':row_number}, ' | '.join(cell.text for cell in row.cells), 'pptx-table'
            elif shape.has_text_frame:
                yield locator, shape.text_frame.text, 'pptx-text'
    for number, slide in enumerate(deck.slides, 1):
        yield from shapes_text(slide.shapes,number)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            yield {'kind':'slide_notes','slide':number}, slide.notes_slide.notes_text_frame.text, 'pptx-notes'


def extract(data, filename):
    ext = validate_container(data, filename)
    units = {'.pdf':pdf_units,'.docx':docx_units,'.pptx':pptx_units}[ext](data)
    return chunks_from_units(units)


def isolated_extract(data, filename):
    """Timeout/memory bounds apply to the parser and its native OCR children."""
    with tempfile.TemporaryDirectory() as temp:
        source = Path(temp)/('source'+Path(filename).suffix.lower())
        source.write_bytes(data)
        result = subprocess.run([sys.executable,str(Path(__file__).resolve()),str(source)],capture_output=True,text=True,timeout=240)
        try:
            payload = json.loads(result.stdout)
        except ValueError as exc:
            raise ValueError('Extraction process stopped or exceeded its resource limit. Split the document and retry.') from exc
        if result.returncode:
            raise ValueError(payload.get('error','Extraction failed.'))
        return payload['chunks']


if __name__ == '__main__':
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS,(1536*1024*1024,1536*1024*1024))
        resource.setrlimit(resource.RLIMIT_CPU,(180,180))
        path = Path(sys.argv[1])
        print(json.dumps({'chunks':extract(path.read_bytes(),path.name)}))
    except Exception as exc:
        print(json.dumps({'error':str(exc)[:1500]}))
        sys.exit(1)
